from datetime import date, datetime
from io import BytesIO
import unittest

import fitz
from openpyxl import load_workbook

from core.models import BankStatement, LedgerSummary, OutsideCashCategory, OutsideCashStatement, ReportSettings, Transaction
from core.reconcile import reconcile
from outputs.bank_report import build_bank_copy_pdf
from outputs.excel_report import build_balance_report_xlsx
from outputs.internal_report import build_internal_check_pdf


class OutputTests(unittest.TestCase):
    def setUp(self):
        ref = datetime(2026, 8, 31)
        self.school = LedgerSummary("school.xlsx", "school", 1500, 1000, 500, ref)
        self.outside = LedgerSummary("outside.xlsx", "outside", 300, 200, 100, ref)
        self.statement = OutsideCashStatement(
            "statement.xlsx", 100,
            categories=[OutsideCashCategory("보관금 / 기타세금", 100, "4대보험 및 기타세금")],
            as_of=ref,
        )
        sb = BankStatement("school.xls", account_number="000-000", account_holder="테스트기관", current_balance=700,
                           transactions=[Transaction(ref.replace(hour=17), withdrawal=20, balance=500, description="대체")])
        ob = BankStatement("outside.xls", account_number="000-001", account_holder="테스트기관", current_balance=120,
                           transactions=[Transaction(ref.replace(hour=9), deposit=10, balance=100, description="대체")])
        cb = BankStatement("card.xls", account_number="000-002", account_holder="테스트기관", current_balance=0,
                           transactions=[
                               Transaction(datetime(2026,8,27,8), deposit=200, balance=200, description="대체"),
                               Transaction(datetime(2026,8,27,17), withdrawal=200, balance=0, description="NH비씨대금", memo="NHBC기업카드"),
                           ])
        self.rec = reconcile([sb,ob,cb], self.school, self.outside, self.statement)

    def test_excel_output_has_expected_sections(self):
        data = build_balance_report_xlsx(self.rec, self.school, self.outside, self.statement, 0)
        wb = load_workbook(BytesIO(data), data_only=True)
        ws = wb["잔액현황"]
        self.assertEqual(ws["A1"].value, "현금 출납부 수입 지출 내역")
        self.assertEqual(ws["E5"].value, 500)
        self.assertEqual(ws["E10"].value, 100)
        self.assertEqual(ws["D14"].value, 200)

    def test_pdf_outputs_render(self):
        settings = ReportSettings(check_date=date(2026,9,15), inspector_title="주무관", inspector_name="테스트", confirmer_title="행정실장", confirmer_name="확인자")
        p1 = build_internal_check_pdf(self.rec, settings)
        p2 = build_bank_copy_pdf(self.rec)
        self.assertTrue(p1.startswith(b"%PDF"))
        self.assertTrue(p2.startswith(b"%PDF"))
        with fitz.open(stream=p1, filetype="pdf") as d1:
            self.assertEqual(d1.page_count, 1)
        with fitz.open(stream=p2, filetype="pdf") as d2:
            self.assertEqual(d2.page_count, 3)


if __name__ == "__main__":
    unittest.main()
