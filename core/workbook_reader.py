from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import BinaryIO

from .models import WorkbookData


def _read_xlsx(data: bytes, filename: str) -> WorkbookData:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheets: dict[str, list[list[object]]] = {}
    try:
        for ws in wb.worksheets:
            rows: list[list[object]] = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
            sheets[ws.title] = rows
    finally:
        wb.close()
    return WorkbookData(filename=filename, sheets=sheets)


def _read_xls(data: bytes, filename: str) -> WorkbookData:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "구형 .xls 파일을 읽으려면 xlrd 패키지가 필요합니다. requirements.txt로 설치해 주세요."
        ) from exc

    book = xlrd.open_workbook(file_contents=data)
    sheets: dict[str, list[list[object]]] = {}
    for sheet in book.sheets():
        rows: list[list[object]] = []
        for r in range(sheet.nrows):
            out: list[object] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value: object = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode))
                    except Exception:
                        value = cell.value
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    value = int(cell.value)
                out.append(value)
            rows.append(out)
        sheets[sheet.name] = rows
    return WorkbookData(filename=filename, sheets=sheets)


def read_workbook(uploaded_file: BinaryIO) -> WorkbookData:
    filename = getattr(uploaded_file, "name", "uploaded")
    data = uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read()
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        return _read_xlsx(data, filename)
    if lower.endswith(".xls"):
        return _read_xls(data, filename)
    raise ValueError(f"지원하지 않는 파일 형식입니다: {filename}")
