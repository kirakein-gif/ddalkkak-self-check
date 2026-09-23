from __future__ import annotations

from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import LedgerSummary, OutsideCashStatement
from core.reconcile import Reconciliation, balance_at_or_before


HEADER_FILL = "D9E1F2"
FONT_NAME = "맑은 고딕"
MONEY_FMT = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"_);_(@_)'
DATE_FMT = "m/d/yy"


def _side(style: str):
    return Side(style=style, color="000000")


def _cell_border(left="thin", right="thin", top="thin", bottom="thin"):
    return Border(left=_side(left), right=_side(right), top=_side(top), bottom=_side(bottom))


def _apply_table_border(ws, min_row, max_row, min_col, max_col, double_header_bottom=None):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            left = "medium" if c == min_col else "thin"
            right = "medium" if c == max_col else "thin"
            top = "medium" if r == min_row else "thin"
            bottom = "medium" if r == max_row else "thin"
            if double_header_bottom and r == double_header_bottom:
                bottom = "double"
            ws.cell(r, c).border = _cell_border(left, right, top, bottom)


def _outside_group_totals(statement: OutsideCashStatement | None) -> tuple[int, int]:
    if not statement:
        return 0, 0
    taxes = sum(x.balance for x in statement.categories if x.group == "4대보험 및 기타세금")
    deposits = sum(x.balance for x in statement.categories if x.group == "계약보증금 및 기타보관금")
    return taxes, deposits


def build_balance_report_xlsx(
    rec: Reconciliation,
    school: LedgerSummary,
    outside: LedgerSummary,
    outside_statement: OutsideCashStatement | None,
    fixed_deposit: int = 0,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "잔액현황"
    ws.sheet_view.showGridLines = False

    widths = {"A": 9, "B": 22, "C": 19, "D": 19, "E": 17, "F": 20, "G": 20, "H": 20}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for r, height in {1: 38, 2: 28, 3: 30, 4: 27, 5: 42, 6: 18, 7: 28, 8: 30, 9: 27, 10: 42, 11: 18, 12: 28, 13: 42}.items():
        ws.row_dimensions[r].height = height

    ws.merge_cells("A1:H1")
    ws["A1"] = "현금 출납부 수입 지출 내역"
    ws["A1"].font = Font(name=FONT_NAME, size=20, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    section_font = Font(name=FONT_NAME, size=12, bold=True)
    for cell, title in [("A2", "1. 학교회계 현금 출납부 수입 지출 내역"), ("A7", "2. 세입세출외 현금 출납부 수입 지출 내역"), ("A12", "3. 법인카드 대금 결제 내역")]:
        ws[cell] = title
        ws[cell].font = section_font
        ws[cell].alignment = Alignment(horizontal="left", vertical="center")

    for rg in ("A3:A4", "B3:B4", "C3:C4", "D3:D4", "E3:E4", "F3:H3"):
        ws.merge_cells(rg)
    headers = {"A3":"회계\n연도", "B3":"기준일", "C3":"수입액", "D3":"지출액", "E3":"잔    액", "F3":"현재 잔액 내역", "F4":"정기예금", "G4":"학교회계통장잔액", "H4":"잔액 합계"}
    for addr, value in headers.items():
        ws[addr] = value

    ref = rec.reference_date or school.as_of or datetime.now()
    school_base = balance_at_or_before(rec.school_bank, rec.reference_date) or 0
    row5 = [ref.year, ref, school.income, school.expense, school.balance, fixed_deposit or None, school_base, school_base + fixed_deposit]
    for c, value in enumerate(row5, 1):
        ws.cell(5, c, value)

    for rg in ("A8:A9", "B8:B9", "C8:C9", "D8:D9", "E8:E9", "F8:H8"):
        ws.merge_cells(rg)
    headers2 = {"A8":"회계\n연도", "B8":"기준일", "C8":"수입액", "D8":"지출액", "E8":"잔    액", "F8":"현재 잔액 내역", "F9":"4대보험 및 기타세금", "G9":"계약보증금 및 기타보관금", "H9":"잔액 합계"}
    for addr, value in headers2.items():
        ws[addr] = value
    taxes, deposits = _outside_group_totals(outside_statement)
    row10 = [ref.year, ref, outside.income, outside.expense, outside.balance, taxes, deposits, taxes + deposits]
    for c, value in enumerate(row10, 1):
        ws.cell(10, c, value)

    ws.merge_cells("F13:H13")
    card_headers = ["회계\n연도", "기준일", "지출액", "대금 결제액", "잔    액"]
    for c, value in enumerate(card_headers, 1):
        ws.cell(13, c, value)
    ws["F13"] = "잔액 내역"

    start = 14
    card_rows = rec.card_checks or []
    for idx, card in enumerate(card_rows, start):
        dt = datetime.strptime(card.date, "%Y-%m-%d")
        vals = [dt.year, dt, card.payment, card.payment, card.balance_after]
        for c, value in enumerate(vals, 1):
            ws.cell(idx, c, value)
        ws.merge_cells(start_row=idx, start_column=6, end_row=idx, end_column=8)
        ws.cell(idx, 6, "")
        ws.row_dimensions[idx].height = 36

    end_card = max(start, start + len(card_rows) - 1)

    for rg in ["A3:H4", "A8:H9", "A13:H13"]:
        for row in ws[rg]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
                cell.font = Font(name=FONT_NAME, size=11, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in [5, 10] + list(range(start, end_card + 1)):
        for c in range(1, 9):
            ws.cell(r, c).font = Font(name=FONT_NAME, size=11)
            ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 2).number_format = DATE_FMT
        for c in range(3, 9):
            ws.cell(r, c).number_format = MONEY_FMT

    _apply_table_border(ws, 3, 5, 1, 8, double_header_bottom=4)
    _apply_table_border(ws, 8, 10, 1, 8, double_header_bottom=9)
    _apply_table_border(ws, 13, end_card, 1, 8, double_header_bottom=13)

    ws.freeze_panes = "A3"
    ws.print_area = f"A1:H{end_card}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def build_reconciliation_xlsx(rec: Reconciliation) -> bytes:
    raise RuntimeError("build_balance_report_xlsx()를 사용하세요.")
